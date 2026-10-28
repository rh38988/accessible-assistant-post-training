r"""CPU-only checks for blinded rating artifacts; never score or generate replies.

Run: .\.venv\Scripts\python.exe -B -m unittest discover -s eval -p "test_*.py" -v
"""
import copy
import csv
import hashlib
import io
import json
import random
import tempfile
import unittest
import zipfile
from collections import Counter
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from eval import make_rating_sheet as rating


class RatingSheetTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.scenario_path = rating.REPO / "eval/scenarios/eval_v1.jsonl"
        cls.outputs_dir = rating.REPO / "eval/outputs"
        cls.scenarios = rating.load_scenarios(cls.scenario_path)
        cls.outputs = rating.load_outputs(cls.outputs_dir)

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def generate(self, name="ratings", seed=42):
        return rating.generate(self.scenario_path, self.outputs_dir, self.root / name, seed)

    def workbook(self, path):
        wb = load_workbook(path)
        self.addCleanup(wb.close)
        return wb

    def test_45_rows_identical_raters_and_matching_private_key_ids(self):
        first, second, key_path = self.generate()
        self.assertTrue(first.read_bytes() == second.read_bytes())
        wb = self.workbook(first)
        ws = wb["Ratings"]
        self.assertEqual(ws.max_row - 1, 45)
        self.assertEqual(tuple(c.value for c in ws[1]), rating.HEADERS)
        items = [dict(zip(rating.HEADERS, row)) for row in ws.iter_rows(min_row=2, values_only=True)]
        with key_path.open(encoding="utf-8", newline="") as stream:
            key = list(csv.DictReader(stream))
        self.assertEqual(len(key), 45)
        self.assertEqual(len({r["item_id"] for r in items}), 45)
        self.assertTrue({r["item_id"] for r in items} == {r["item_id"] for r in key})
        self.assertEqual(Counter(r["scenario_id"] for r in items),
                         Counter({s["id"]: 3 for s in self.scenarios}))
        lookup = {r["item_id"]: r for r in key}
        scenarios = {s["id"]: s for s in self.scenarios}
        outputs = {c: {r["scenario_id"]: r for r in rows} for c, rows in self.outputs.items()}
        for item in items:
            private = lookup[item["item_id"]]
            self.assertTrue(item["scenario_id"] == private["scenario_id"])
            self.assertTrue(item["candidate"] == private["candidate"])
            # Boolean assertions keep private mappings out of failure output.
            expected = outputs[private["condition"]][item["scenario_id"]]["response"]
            self.assertTrue(item["response"] == expected)
            sc = scenarios[item["scenario_id"]]
            self.assertEqual(item["prompt"], "\n\n".join(
                f"{t['role'].upper()}: {t['content']}" for t in sc["prompt"]))
            for field in ("clarification_required", "risk_level", "source_url", "task_family"):
                self.assertEqual(item[field], sc[field])
            self.assertEqual(item["expected_elements"], "\n".join(sc["expected_elements"]))
            self.assertEqual(item["word_count"], len(item["response"].split()))
            self.assertTrue(all(item[field] is None for field in (*rating.SCORES, *rating.CHECKS, "notes")))
        for sc in self.scenarios:
            group = [row for row in key if row["scenario_id"] == sc["id"]]
            self.assertTrue({r["candidate"] for r in group} == set(rating.CANDIDATES))
            self.assertTrue({r["condition"] for r in group} == set(rating.CONDITIONS))

    def test_seed_42_is_deterministic_independent_per_scenario_and_shuffles_rows(self):
        first, key = rating.build_items(self.scenarios, self.outputs)
        again, again_key = rating.build_items(self.scenarios, self.outputs, 42)
        self.assertTrue(first == again and key == again_key)
        reversed_outputs = {c: list(reversed(rows)) for c, rows in self.outputs.items()}
        reordered, reordered_key = rating.build_items(list(reversed(self.scenarios)), reversed_outputs, 42)
        self.assertTrue(first == reordered and key == reordered_key)
        changed, changed_key = rating.build_items(self.scenarios, self.outputs, 43)
        self.assertTrue(first != changed and key != changed_key)
        rng = random.Random(42)
        expected = []
        permutations = set()
        for sc in sorted(self.scenarios, key=lambda sc: sc["id"]):
            conditions = list(rating.CONDITIONS)
            rng.shuffle(conditions)
            permutations.add(tuple(conditions))
            expected.extend((sc["id"], label, condition)
                            for label, condition in zip(rating.CANDIDATES, conditions))
        self.assertTrue([(k["scenario_id"], k["candidate"], k["condition"]) for k in key] == expected)
        self.assertGreater(len(permutations), 1)
        ordered_ids = [k["item_id"] for k in key]
        rng.shuffle(ordered_ids)
        self.assertEqual([item["item_id"] for item in first], ordered_ids)
        self.assertTrue(ordered_ids != [k["item_id"] for k in key])
        a = self.workbook(self.generate("first")[0])
        b = self.workbook(self.generate("second")[0])
        self.assertTrue(list(a["Ratings"].values) == list(b["Ratings"].values))

    def test_no_condition_leakage_in_cells_metadata_or_hidden_parts(self):
        outputs = copy.deepcopy(self.outputs)
        for rows in outputs.values():
            for row in rows:
                row["adapter"] = "PRIVATE_ADAPTER_METADATA_DO_NOT_EXPORT"
                row["response"] = "A candidate answer without identity metadata."
        selection = rating.select_checkpoint_scenarios(self.scenarios, 42)
        items, _ = rating.build_items(self.scenarios, outputs, 42, [sc["id"] for sc in selection])
        wb = rating.make_workbook(items, rating.REPO / "eval/rubric_v1.md", selection, 42)
        path = self.root / "anonymous.xlsx"
        wb.save(path)
        wb.close()
        loaded = self.workbook(path)
        self.assertEqual(loaded.sheetnames, ["Ratings", "Rubric", "Selection"])
        self.assertFalse(loaded.defined_names)
        self.assertFalse(loaded.properties.creator)
        self.assertFalse(loaded.properties.lastModifiedBy)
        for ws in loaded:
            self.assertEqual(ws.sheet_state, "visible")
            for row in ws:
                for cell in row:
                    self.assertIsNone(cell.comment)
                    self.assertIsNone(cell.hyperlink)
                    self.assertNotEqual(cell.data_type, "f")
        self.assertTrue(all(set(item) == set(rating.HEADERS) for item in items))
        self.assertFalse(set(rating.HEADERS) & {"condition", "adapter", "model", "output_file"})
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                data = archive.read(name)
                for marker in (b"PRIVATE_ADAPTER_METADATA", b"eval_v1-original.jsonl",
                               b"eval_v1-prompt.jsonl", b"eval_v1-sft.jsonl", b"key.csv"):
                    self.assertNotIn(marker, data)
                self.assertNotIn("hidden", name.lower())

    def test_exact_rubric_and_score_check_choices_and_applicability(self):
        first, _, _ = self.generate()
        wb = self.workbook(first)
        expected = (rating.REPO / "eval/rubric_v1.md").read_text(encoding="utf-8").splitlines()
        actual = [row[0] or "" for row in wb["Rubric"].iter_rows(values_only=True)]
        self.assertEqual(actual[:len(expected)], expected)
        ws = wb["Ratings"]
        validations = ws.data_validations.dataValidation
        self.assertEqual(len(validations), 7)
        for field, formula in [(f, '"0,1,2,NA"') for f in rating.SCORES] + [
            ("clarification_behavior", '"appropriate,guessed,over-asked"'),
            ("unsafe_advice", '"yes,no"'), ("language_match", '"yes,no"')
        ]:
            col = rating.HEADERS.index(field) + 1
            for row in range(2, 47):
                item = dict(zip(rating.HEADERS, (c.value for c in ws[row])))
                applicable = (field != "unsafe_advice" or item["risk_level"] in {"medium", "high"})
                applicable &= field != "language_match" or item["language"] == "zh-TW"
                rules = [rule for rule in validations if ws.cell(row, col).coordinate in rule]
                self.assertEqual(len(rules), int(applicable))
                if rules:
                    self.assertEqual(rules[0].formula1, formula)
                    self.assertTrue(rules[0].allow_blank)

    def test_rejects_incomplete_duplicate_mismatched_or_empty_outputs(self):
        for case in ("missing", "duplicate", "unknown", "condition", "empty"):
            outputs = copy.deepcopy(self.outputs)
            rows = outputs["original"]
            if case == "missing":
                rows.pop()
            elif case == "duplicate":
                rows[-1] = copy.deepcopy(rows[0])
            elif case == "unknown":
                rows[0]["scenario_id"] = "UNKNOWN"
            elif case == "condition":
                rows[0]["condition"] = "other"
            else:
                rows[0]["response"] = " "
            with self.subTest(case=case), self.assertRaises(ValueError):
                rating.build_items(self.scenarios, outputs)

    def test_only_three_named_files_are_read(self):
        directory = self.root / "outputs"
        directory.mkdir()
        for condition, rows in self.outputs.items():
            (directory / f"eval_v1-{condition}.jsonl").write_text(
                "\n".join(json.dumps(row) for row in rows), encoding="utf-8")
        (directory / "unrelated.jsonl").write_text("not JSON", encoding="utf-8")
        self.assertTrue(rating.load_outputs(directory) == self.outputs)

    def test_inputs_unchanged_existing_ratings_protected_and_cli_keeps_key_private(self):
        sources = [self.scenario_path, rating.REPO / "eval/rubric_v1.md"] + [
            self.outputs_dir / f"eval_v1-{condition}.jsonl" for condition in rating.CONDITIONS]
        before = {path: hashlib.sha256(path.read_bytes()).digest() for path in sources}
        with patch("sys.stdout", new_callable=io.StringIO) as stdout:
            rating.main(["--ratings-dir", str(self.root / "cli"), "--seed", "42"])
        self.assertNotIn("original", stdout.getvalue())
        self.assertNotIn("sft", stdout.getvalue())
        self.assertTrue(before == {path: hashlib.sha256(path.read_bytes()).digest() for path in sources})
        artifacts = [self.root / "cli" / name for name in rating.FILENAMES]
        hashes = [hashlib.sha256(path.read_bytes()).digest() for path in artifacts]
        with self.assertRaises(ValueError):
            self.generate("cli")
        self.assertTrue(hashes == [hashlib.sha256(path.read_bytes()).digest() for path in artifacts])

    def test_formula_like_response_is_preserved_as_text(self):
        items, _ = rating.build_items(self.scenarios, self.outputs)
        items[0]["response"] = '=HYPERLINK("https://example.invalid", "text")'
        wb = rating.make_workbook(items, rating.REPO / "eval/rubric_v1.md")
        path = self.root / "literal.xlsx"
        wb.save(path)
        wb.close()
        loaded = self.workbook(path)
        cell = loaded["Ratings"].cell(2, rating.HEADERS.index("response") + 1)
        self.assertEqual(cell.value, items[0]["response"])
        self.assertEqual(cell.data_type, "s")

    def test_five_selected_scenarios_have_all_required_metadata_coverage(self):
        selected = rating.select_checkpoint_scenarios(self.scenarios, 42)
        self.assertEqual(len(selected), 5)
        self.assertEqual(len({sc["id"] for sc in selected}), 5)
        counts = Counter(sc["task_family"] for sc in selected)
        self.assertEqual(set(counts), {sc["task_family"] for sc in self.scenarios})
        self.assertEqual(sorted(counts.values()), [1, 2, 2])
        self.assertTrue(any(sc["id"] in rating.ZH_TW_IDS for sc in selected))
        self.assertTrue(any(len(sc["prompt"]) > 1 for sc in selected))
        self.assertTrue(any(sc["clarification_required"] for sc in selected))
        self.assertTrue(any(sc["task_family"] == "Online Safety and Scam Prevention"
                            and sc["risk_level"] in {"medium", "high"} for sc in selected))

    def test_checkpoint_selection_is_deterministic_and_ignores_text_and_input_order(self):
        def ids(scenarios, seed=42):
            return [sc["id"] for sc in rating.select_checkpoint_scenarios(scenarios, seed)]

        expected = ids(self.scenarios)
        self.assertEqual(ids(self.scenarios), expected)
        self.assertEqual(ids(list(reversed(self.scenarios))), expected)
        changed = copy.deepcopy(self.scenarios)
        for sc in changed:
            for turn in sc["prompt"]:
                turn["content"] = "Arbitrary text that must not affect selection."
            sc["expected_elements"] = ["Ignored"]
            sc["source_url"] = "https://example.invalid/ignored"
            sc["response"] = "Injected response content must never affect selection."
            sc["adapter"] = "Injected adapter must never affect selection."
        self.assertEqual(ids(changed), expected)
        self.assertNotEqual(ids(self.scenarios, 43), expected)
        impossible = copy.deepcopy(self.scenarios)
        for sc in impossible:
            sc["clarification_required"] = False
        with self.assertRaises(ValueError):
            rating.select_checkpoint_scenarios(impossible)

    def test_pilot_has_15_rows_identical_workbooks_matching_ids_and_recorded_seed(self):
        paths = rating.generate(self.scenario_path, self.outputs_dir, self.root / "pilot", 42, pilot=True)
        first, second, key_path = paths
        self.assertTrue(first.read_bytes() == second.read_bytes())
        wb = self.workbook(first)
        ws = wb["Ratings"]
        self.assertEqual(ws.max_row - 1, 15)
        rows = [dict(zip(rating.HEADERS, row)) for row in ws.iter_rows(min_row=2, values_only=True)]
        selected = rating.select_checkpoint_scenarios(self.scenarios, 42)
        selected_ids = {sc["id"] for sc in selected}
        self.assertEqual(Counter(row["scenario_id"] for row in rows), Counter({sid: 3 for sid in selected_ids}))
        for sid in selected_ids:
            self.assertEqual({row["candidate"] for row in rows if row["scenario_id"] == sid}, {"A", "B", "C"})
        self.assertTrue(all(row[field] is None for row in rows for field in (*rating.SCORES, *rating.CHECKS, "notes")))
        with key_path.open(encoding="utf-8", newline="") as stream:
            key = list(csv.DictReader(stream))
        self.assertEqual(len(key), 15)
        self.assertTrue({row["item_id"] for row in key} == {row["item_id"] for row in rows})
        self.assertTrue(all({row["condition"] for row in key if row["scenario_id"] == sid}
                            == set(rating.CONDITIONS) for sid in selected_ids))
        lookup = {row["item_id"]: row for row in key}
        outputs = {c: {row["scenario_id"]: row for row in values} for c, values in self.outputs.items()}
        for row in rows:
            private = lookup[row["item_id"]]
            self.assertTrue(row["response"] == outputs[private["condition"]][row["scenario_id"]]["response"])
        record = list(wb["Selection"].values)
        self.assertEqual(record[0], ("selector", "metadata-combinations-v1"))
        self.assertEqual(record[1], ("selection_seed", 42))
        self.assertEqual(record[2], ("randomization_seed", 42))
        self.assertEqual({row[0] for row in record[5:]}, selected_ids)
        again = rating.generate(self.scenario_path, self.outputs_dir, self.root / "again", 42, pilot=True)
        other = self.workbook(again[0])
        self.assertTrue(list(ws.values) == list(other["Ratings"].values))
        self.assertEqual(record, list(other["Selection"].values))

    def test_replace_blank_45_rows_with_pilot_only_overwrites_three_artifacts(self):
        paths = self.generate("replace")
        directory = paths[0].parent
        unrelated = directory / "leave_me.txt"
        unrelated.write_text("preserved", encoding="utf-8")
        before_names = {path.name for path in directory.iterdir()}
        sources = [self.scenario_path, rating.REPO / "eval/rubric_v1.md"] + [
            self.outputs_dir / f"eval_v1-{condition}.jsonl" for condition in rating.CONDITIONS]
        hashes = {p: hashlib.sha256(p.read_bytes()).digest() for p in sources}
        with patch("sys.stdout", new_callable=io.StringIO) as stdout:
            rating.main(["--ratings-dir", str(directory), "--pilot", "--replace-blank", "--seed", "42"])
        self.assertIn("15 rating rows each", stdout.getvalue())
        self.assertIn("seed: 42", stdout.getvalue())
        for sc in rating.select_checkpoint_scenarios(self.scenarios, 42):
            self.assertIn(sc["id"], stdout.getvalue())
        self.assertNotIn("original", stdout.getvalue())
        self.assertNotIn("sft", stdout.getvalue())
        self.assertEqual(self.workbook(paths[0])["Ratings"].max_row - 1, 15)
        self.assertTrue(paths[0].read_bytes() == paths[1].read_bytes())
        self.assertEqual({path.name for path in directory.iterdir()}, before_names)
        self.assertEqual(unrelated.read_text(encoding="utf-8"), "preserved")
        self.assertTrue(hashes == {p: hashlib.sha256(p.read_bytes()).digest() for p in sources})

    def test_replace_aborts_without_changing_any_artifact_if_either_rater_has_any_entry(self):
        paths = self.generate("guard")
        originals = [path.read_bytes() for path in paths]
        fields = (*rating.SCORES, *rating.CHECKS, "notes")
        for rater in (0, 1):
            for field in fields:
                with self.subTest(rater=rater + 1, field=field):
                    wb = load_workbook(io.BytesIO(originals[rater]))
                    # Includes numeric zero; truthiness must not treat it as blank.
                    wb["Ratings"].cell(2, rating.HEADERS.index(field) + 1, 0)
                    wb.save(paths[rater])
                    wb.close()
                    before = [hashlib.sha256(path.read_bytes()).digest() for path in paths]
                    with self.assertRaisesRegex(ValueError, "entered ratings or notes"):
                        rating.generate(self.scenario_path, self.outputs_dir, paths[0].parent,
                                        42, pilot=True, replace_blank=True)
                    self.assertTrue(before == [hashlib.sha256(path.read_bytes()).digest() for path in paths])
                    paths[rater].write_bytes(originals[rater])

    def test_replace_guard_rejects_formulas_missing_fields_and_missing_files(self):
        paths = self.generate("fail_closed")
        original = paths[1].read_bytes()
        for case in ("formula", "missing_fields"):
            with self.subTest(case=case):
                wb = load_workbook(io.BytesIO(original))
                cell = "M2" if case == "formula" else "M1"
                wb["Ratings"][cell] = '=IF(1=1,"",0)' if case == "formula" else "unknown"
                wb.save(paths[1])
                wb.close()
                before = [hashlib.sha256(path.read_bytes()).digest() for path in paths]
                with self.assertRaises(ValueError):
                    rating.generate(self.scenario_path, self.outputs_dir, paths[0].parent,
                                    42, pilot=True, replace_blank=True)
                self.assertTrue(before == [hashlib.sha256(path.read_bytes()).digest() for path in paths])
        with self.assertRaises(ValueError):
            rating.generate(self.scenario_path, self.outputs_dir, self.root / "missing",
                            42, pilot=True, replace_blank=True)


if __name__ == "__main__":
    unittest.main()
